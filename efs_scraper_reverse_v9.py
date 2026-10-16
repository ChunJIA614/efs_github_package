from __future__ import annotations

import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

def _ensure_dependencies() -> None:
    """Auto-install required third-party packages if missing on a personal PC."""
    missing = []
    for mod_name, pkg_name in (
        ("openpyxl", "openpyxl"),
        ("bs4", "beautifulsoup4"),
        ("playwright", "playwright"),
    ):
        try:
            __import__(mod_name)
        except ImportError:
            missing.append(pkg_name)
    if missing:
        print(f"[Setup] Installing missing packages on this PC: {', '.join(missing)} ...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", *missing])


try:
    _ensure_dependencies()
    import openpyxl
    from bs4 import BeautifulSoup
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from playwright.sync_api import Page, Response, TimeoutError as PlaywrightTimeoutError, sync_playwright
except Exception as _boot_exc:
    import traceback
    print("\n" + "=" * 76)
    print("  [STARTUP ERROR] Failed to load or install required Python packages:")
    print("=" * 76)
    traceback.print_exc()
    try:
        input("\nPress ENTER to close this window...")
    except Exception:
        time.sleep(15)
    sys.exit(1)

APP_TITLE = "Malaysia Court E-Filing (EFS) Scraper — Final Edition"
URL = "https://efs.kehakiman.gov.my/EFSWeb/eFiling/MyFiling.aspx"
BASE_DIR = Path(__file__).resolve().parent

# Keep OUTPUT_DIR in the script's folder (e.g. \\server\Public\ELLA\EFS\efs_output),
# but store PROFILE_DIR on the local machine (%LOCALAPPDATA%\EFS_Browser_Profile)
# so Chromium never fails with ProcessSingleton / Error code 32 on network shares!
LOCAL_APP_DIR = Path(os.environ.get("LOCALAPPDATA") or tempfile.gettempdir())
PROFILE_DIR = LOCAL_APP_DIR / "EFS_Browser_Profile"
OUTPUT_DIR = BASE_DIR / "efs_output"

GRID_ID = "ctl00_ctl00_Body_MyFilingBody_radGridMyFiling"
GRID_TABLE_ID = f"{GRID_ID}_ctl00"
CASE_GRID_PREFIX = "ctl00_ctl00_Body_MyFilingBody_CaseUserControl_radGridCaseDetails"
CASE_GRID_TABLE_ID = f"{CASE_GRID_PREFIX}_ctl00"
ROWS_SELECTOR = (
    f"#{GRID_TABLE_ID} tbody tr.rgRow, "
    f"#{GRID_TABLE_ID} tbody tr.rgAltRow"
)
PREVIOUS_BUTTON_SELECTOR = f"#{GRID_TABLE_ID} input.rgPagePrev"
NEXT_BUTTON_SELECTOR = f"#{GRID_TABLE_ID} input.rgPageNext"
PAGE_SIZE_INPUT = f"#{GRID_TABLE_ID}_ctl03_ctl01_PageSizeComboBox_Input"

OUTPUT_HEADERS = ["DATE", "NAME", "PRODUCT", "COURT", "PRICE", "TOTAL"]

# Matches both comma-formatted (1,000.00) and plain (1000.00) decimal amounts
MONEY_RE = re.compile(r"(?<!\d)(\d+(?:,\d{3})*\.\d{2})(?!\d)")


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def parse_money_decimal(text: str) -> Decimal | None:
    match = MONEY_RE.search(norm(text))
    if not match:
        return None
    try:
        return Decimal(match.group(1).replace(",", ""))
    except InvalidOperation:
        return None


def parse_user_date(s: str) -> date | None:
    """Parse a date string in DD/MM/YYYY, YYYY-MM-DD, DD-MM-YYYY, or DD MMM YYYY format."""
    cleaned = norm(s)
    if not cleaned:
        return None
    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d.%m.%Y",
        "%d %b %Y",
        "%d %B %Y",
        "%d/%m/%y",
        "%d-%m-%y",
    ):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def normalize_date_range(
    start_str: str,
    end_str: str,
) -> tuple[date | None, date | None, str | None]:
    """Validate and normalize user-supplied From/To date strings.

    Returns (start_date, end_date, error_message).
    If both dates are provided and start_date > end_date, they are automatically
    swapped so start_date is always the earlier date and end_date is the later date.
    """
    s_clean = norm(start_str)
    e_clean = norm(end_str)

    start_dt = parse_user_date(s_clean) if s_clean else None
    if s_clean and start_dt is None:
        return None, None, f"Invalid From Date '{s_clean}'. Please use DD/MM/YYYY (e.g. 05/10/2026)."

    end_dt = parse_user_date(e_clean) if e_clean else None
    if e_clean and end_dt is None:
        return None, None, f"Invalid To Date '{e_clean}'. Please use DD/MM/YYYY (e.g. 14/10/2026)."

    if start_dt and end_dt and start_dt > end_dt:
        start_dt, end_dt = end_dt, start_dt

    return start_dt, end_dt, None


def is_date_in_range(
    row_date_str: str,
    start_dt: date | None,
    end_dt: date | None,
) -> bool:
    """Check whether a row's DD/MM/YYYY filing date falls within [start_dt, end_dt]."""
    if start_dt is None and end_dt is None:
        return True
    row_dt = parse_user_date(row_date_str)
    if row_dt is None:
        return True
    if start_dt is not None and row_dt < start_dt:
        return False
    if end_dt is not None and row_dt > end_dt:
        return False
    return True


def clean_stale_profile_locks(profile_dir: Path) -> None:
    """Remove stale Chromium SingletonLock files if they exist and are not actively locked."""
    if not profile_dir.exists():
        return
    for lock_name in ("SingletonLock", "SingletonCookie", "SingletonSocket", "lockfile"):
        lock_file = profile_dir / lock_name
        if lock_file.exists():
            try:
                lock_file.unlink()
            except Exception:
                pass


def launch_edge_resilient(playwright_instance: Any, log_fn: Any = print) -> tuple[Any, Page, Path | None]:
    """Launch Microsoft Edge resiliently, avoiding ProcessSingleton / UNC network share lock errors.

    Tries:
      1. Local persistent profile (%LOCALAPPDATA%\\EFS_Browser_Profile)
      2. Fresh session-specific local profile if another Edge instance holds the lock
      3. Standard non-persistent Edge context as ultimate fallback
    """
    PROFILE_DIR.mkdir(parents=True, exist_ok=True)
    clean_stale_profile_locks(PROFILE_DIR)

    # Attempt 1: Local persistent profile
    try:
        log_fn(f"Launching Microsoft Edge (Local Profile: {PROFILE_DIR})...")
        context = playwright_instance.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            channel="msedge",
            headless=False,
            viewport={"width": 1440, "height": 900},
        )
        page = context.pages[0] if context.pages else context.new_page()
        return context, page, None
    except Exception as first_exc:
        log_fn(f"Primary profile busy ({first_exc.__class__.__name__}); switching to isolated local session profile...")

    # Attempt 2: Isolated session profile on local disk
    temp_profile = Path(tempfile.mkdtemp(prefix="efs_edge_session_", dir=str(LOCAL_APP_DIR)))
    try:
        context = playwright_instance.chromium.launch_persistent_context(
            user_data_dir=str(temp_profile),
            channel="msedge",
            headless=False,
            viewport={"width": 1440, "height": 900},
        )
        page = context.pages[0] if context.pages else context.new_page()
        return context, page, temp_profile
    except Exception:
        shutil.rmtree(temp_profile, ignore_errors=True)

    # Attempt 3: Standard browser launch without user-data-dir lock
    browser = playwright_instance.chromium.launch(
        channel="msedge",
        headless=False,
    )
    context = browser.new_context(viewport={"width": 1440, "height": 900})
    page = context.new_page()
    return context, page, None


def extract_filing_date(status_text: str) -> str:
    m = re.search(
        r"\b(\d{1,2})\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\s+(\d{4})\b",
        status_text or "",
        flags=re.I,
    )
    if not m:
        return ""
    try:
        return datetime.strptime(
            f"{m.group(1)} {m.group(2)} {m.group(3)}", "%d %b %Y"
        ).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def clean_product(s: str) -> str:
    s = norm(s)
    # Remove uploaded PDF filename (including filenames with nested parentheses like (S16).pdf)
    s = re.sub(r"\s*\((?:[^()]|\([^()]*\))*\.pdf\)\s*$", "", s, flags=re.I)
    s = re.sub(r"\s*<br\s*/?>\s*$", "", s, flags=re.I)
    return s.strip(" -:|")


def extract_person_from_product_text(raw_product: str) -> str:
    """Extract the party name embedded in an EFS PDF filename when present."""
    value = norm(raw_product)
    m = re.search(r"\(([^()]*)\.pdf\)\s*$", value, flags=re.I)
    if not m:
        return ""

    inside = norm(m.group(1))
    parts = re.split(r"\s+[-–—]\s+", inside, maxsplit=1)
    candidate = parts[0].strip()
    if not re.search(r"\b(?:BIN|BINTI|A/L|A/P)\b", candidate, flags=re.I):
        return ""
    return clean_person_name(candidate)


def court_code(court: str) -> str:
    return norm(court)


def clean_person_name(name: str) -> str:
    """Clean and normalize person name extracted from EFS.

    Preserves legitimate names starting with 'N' (such as NORFITRI, NOOR, NURUL,
    NIK, NAZRI, NG), while stripping extraneous prefixes/suffixes such as
    'dan', 'and', 'dan lain-lain', and bracketed IC/ID numbers.
    """
    value = norm(name).upper()
    value = re.sub(r"^(?:DAN|AND)\s+", "", value, flags=re.I).strip()
    value = re.sub(
        r"\s+(?:DAN\s+LAIN-LAIN|DAN\s+SEORANG\s+LAGI|&\s+ANOR|&\s+ORS|ET\s+AL)\b.*$",
        "",
        value,
        flags=re.I,
    ).strip()
    value = re.sub(r"\s*\(.*?\)", "", value).strip()
    return value.strip(" -:,;/.")


def find_person_name(text: str) -> str:
    """Extract an individual person's name from party details text.

    Filters out corporate entities (banks, companies with BERHAD/BHD/SDN BHD)
    and extracts the individual party name.
    """
    text = text or ""
    lines = [norm(l) for l in text.splitlines() if norm(l)]

    # Pass 1: Look for individual person lines with IC/ID, skipping corporate entities
    for line in lines:
        line_clean = re.sub(r"^(?:dan|and)\s+", "", line, flags=re.I).strip()
        line_clean = re.sub(
            r"\s+(?:dan\s+lain-lain|dan\s+seorang\s+lagi|&\s+anor|&\s+ors|et\s+al)\b.*$",
            "",
            line_clean,
            flags=re.I,
        ).strip()

        if re.search(r"\b(?:BERHAD|BHD|SDN|BANK|KOPERASI|LEMBAGA|MAJLIS)\b", line_clean, re.I):
            continue

        m = re.search(r"^([A-Za-z][A-Za-z.\s\x27\-\/]+?)\s*\(\s*\d{6,}", line_clean)
        if m:
            cand = clean_person_name(m.group(1))
            if len(cand) > 2:
                return cand

        if re.search(r"\b(?:BIN|BINTI|A/L|A/P)\b", line_clean, re.I):
            cand = clean_person_name(line_clean)
            if len(cand) > 4:
                return cand

    # Pass 2: Any line with (ID), even if corporate filter matched
    for line in lines:
        line_clean = re.sub(r"^(?:dan|and)\s+", "", line, flags=re.I).strip()
        line_clean = re.sub(
            r"\s+(?:dan\s+lain-lain|dan\s+seorang\s+lagi|&\s+anor|&\s+ors|et\s+al)\b.*$",
            "",
            line_clean,
            flags=re.I,
        ).strip()
        m = re.search(r"^([A-Za-z][A-Za-z.\s\x27\-\/]+?)\s*\(\s*\d{6,}", line_clean)
        if m:
            cand = clean_person_name(m.group(1))
            if len(cand) > 2:
                return cand

    # Pass 3: General regex fallback
    pat = re.compile(
        r"\b([A-Z][A-Z.\-']+(?:\s+[A-Z][A-Z.\-']+){1,8})"
        r"\s*\(\s*\d{6,}\s*\)",
        flags=re.I,
    )
    for m in pat.finditer(text):
        cand = clean_person_name(m.group(1))
        if cand and not re.search(r"\b(?:BERHAD|BHD|SDN|BANK)\b", cand, re.I):
            return cand

    return ""


def page_rows(page: Page) -> list[dict[str, str]]:
    rows = page.locator(ROWS_SELECTOR)
    out: list[dict[str, str]] = []

    for i in range(rows.count()):
        row = rows.nth(i)
        vals = [norm(x) for x in row.locator(":scope > td").all_inner_texts()]
        if len(vals) < 6:
            continue
        out.append(
            {
                "filing_id": vals[0],
                "court_full": vals[1],
                "total_documents": vals[2],
                "extraction_code": vals[3],
                "total_amount": vals[4],
                "status": vals[5],
                "filing_date": extract_filing_date(vals[5]),
                "row_index": str(i),
            }
        )
    return out


def get_grid_client_state(page: Page) -> dict[str, Any]:
    try:
        return page.evaluate(
            """() => {
                const id = "ctl00_ctl00_Body_MyFilingBody_radGridMyFiling";
                const g = window.$find ? window.$find(id) : null;
                if (!g) return {ok:false};
                const t = g.get_masterTableView();
                return {
                    ok: true,
                    pageIndex: t.get_currentPageIndex(),
                    pageSize: t.get_pageSize(),
                    pageCount: t.get_pageCount(),
                    virtualItemCount: t.get_virtualItemCount()
                };
            }"""
        )
    except Exception:
        return {"ok": False}


def current_page_number(page: Page) -> int:
    state = get_grid_client_state(page)
    if state.get("ok") and isinstance(state.get("pageIndex"), int):
        return int(state["pageIndex"]) + 1

    try:
        spans = page.locator(f"#{GRID_TABLE_ID} .rgCurrentPage span")
        for i in range(spans.count()):
            txt = norm(spans.nth(i).inner_text())
            if txt.isdigit():
                return int(txt)
    except Exception:
        pass
    return 0


def current_page_size(page: Page) -> str:
    try:
        return page.locator(PAGE_SIZE_INPUT).input_value(timeout=1500)
    except Exception:
        state = get_grid_client_state(page)
        if state.get("ok"):
            return str(state.get("pageSize", ""))
        return ""


def parse_ajax_delta(response_text: str) -> list[dict[str, str]]:
    segments: list[dict[str, str]] = []
    pos = 0
    n = len(response_text)

    while pos < n:
        sep = response_text.find("|", pos)
        if sep < 0:
            break
        length_text = response_text[pos:sep]
        if not length_text.isdigit():
            next_sep = response_text.find("|", sep + 1)
            if next_sep < 0:
                break
            pos = next_sep + 1
            continue

        length = int(length_text)
        pos = sep + 1

        type_end = response_text.find("|", pos)
        if type_end < 0:
            break
        seg_type = response_text[pos:type_end]
        pos = type_end + 1

        id_end = response_text.find("|", pos)
        if id_end < 0:
            break
        seg_id = response_text[pos:id_end]
        pos = id_end + 1

        content = response_text[pos:pos + length]
        segments.append({"type": seg_type, "id": seg_id, "content": content})
        pos += length

        if pos < n and response_text[pos] == "|":
            pos += 1

    return segments


def find_case_panel_html(delta_text: str) -> str:
    segments = parse_ajax_delta(delta_text)
    preferred: list[str] = []
    fallback: list[str] = []

    for seg in segments:
        if seg.get("type") != "updatePanel":
            continue
        seg_id = seg.get("id", "")
        content = seg.get("content", "")
        if "CaseUserControl" in seg_id or "radGridCaseDetails" in content:
            preferred.append(content)
        elif "panelCaseDetails" in seg_id:
            fallback.append(content)

    if preferred:
        return "\n".join(preferred)
    return "\n".join(fallback)


def extract_name_from_case_html(case_html: str) -> str:
    if not case_html:
        return ""
    soup = BeautifulSoup(case_html, "html.parser")
    text = soup.get_text("\n", strip=True)
    return find_person_name(text)


def compute_effective_row_price(
    vals: list[str],
    price_idx: int,
    deposit_idx: int,
    jtt_idx: int,
) -> str | None:
    """Calculate the effective document price across Fi Pemfailan, Amaun Deposit, and JTT columns.

    In EFS:
    - Standard filings place the fee in 'Fi Pemfailan (RM)'.
    - Deposit filings set 'Fi Pemfailan (RM)' to 0.00 and place the amount in 'Amaun Deposit'.
    """
    fee_dec = parse_money_decimal(vals[price_idx]) if 0 <= price_idx < len(vals) else None
    dep_dec = parse_money_decimal(vals[deposit_idx]) if 0 <= deposit_idx < len(vals) else None
    jtt_dec = parse_money_decimal(vals[jtt_idx]) if 0 <= jtt_idx < len(vals) else None

    if fee_dec is None and dep_dec is None and jtt_dec is None:
        return None

    effective = Decimal("0.00")
    if fee_dec is not None:
        effective += fee_dec
    if dep_dec is not None and dep_dec > Decimal("0.00"):
        effective += dep_dec
    elif effective == Decimal("0.00") and jtt_dec is not None and jtt_dec > Decimal("0.00"):
        effective += jtt_dec

    return f"{effective:.2f}"


def extract_document_rows_from_case_html(case_html: str) -> list[dict[str, str]]:
    """Extract actual document rows and their corresponding case party names from CaseUserControl HTML."""
    if not case_html:
        return []

    soup = BeautifulSoup(case_html, "html.parser")
    results: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    detail_tables = soup.find_all("table", class_=lambda c: c and "rgDetailTable" in c)
    if not detail_tables:
        for candidate in soup.find_all("table"):
            thead = candidate.find("thead", recursive=False)
            if thead:
                h_text = norm(thead.get_text(" ", strip=True)).lower()
                if "jenis dokumen" in h_text and ("fi pemfailan" in h_text or "deposit" in h_text):
                    detail_tables.append(candidate)

    fallback_name = find_person_name(soup.get_text("\n", strip=True))

    for table in detail_tables:
        header_cells = []
        thead = table.find("thead", recursive=False)
        if thead:
            for tr in thead.find_all("tr", recursive=False):
                cells = tr.find_all(["th", "td"], recursive=False)
                texts = [norm(c.get_text(" ", strip=True)) for c in cells]
                low = [x.lower() for x in texts]
                if any("jenis dokumen" in x for x in low) and (
                    any("fi pemfailan" in x for x in low) or any("deposit" in x for x in low)
                ):
                    header_cells = texts
                    break
        if not header_cells:
            for tr in table.find_all("tr", recursive=False):
                cells = tr.find_all(["th", "td"], recursive=False)
                texts = [norm(c.get_text(" ", strip=True)) for c in cells]
                low = [x.lower() for x in texts]
                if any("jenis dokumen" in x for x in low) and (
                    any("fi pemfailan" in x for x in low) or any("deposit" in x for x in low)
                ):
                    header_cells = texts
                    break
        if not header_cells:
            continue

        low_header = [x.lower() for x in header_cells]
        product_idx = next((i for i, v in enumerate(low_header) if "jenis dokumen" in v), -1)
        price_idx = next((i for i, v in enumerate(low_header) if "fi pemfailan" in v), -1)
        deposit_idx = next((i for i, v in enumerate(low_header) if "deposit" in v), -1)
        jtt_idx = next((i for i, v in enumerate(low_header) if "jtt" in v), -1)
        if product_idx < 0 or (price_idx < 0 and deposit_idx < 0):
            continue

        # Locate the preceding master case row to extract this case's specific party name
        case_party_name = ""
        detail_tr = table.find_parent("tr")
        if detail_tr:
            curr = detail_tr
            while curr:
                curr = curr.find_previous_sibling("tr")
                if curr and any(c in curr.get("class", []) for c in ["rgRow", "rgAltRow"]):
                    tds = curr.find_all("td", recursive=False)
                    if len(tds) >= 3:
                        party_text = tds[2].get_text("\n", strip=True)
                        case_party_name = find_person_name(party_text)
                    break

        if not case_party_name:
            case_party_name = fallback_name

        min_required_idx = max(product_idx, price_idx, deposit_idx)
        tbody = table.find("tbody", recursive=False) or table
        for tr in tbody.find_all("tr", recursive=False):
            cells = tr.find_all("td", recursive=False)
            if len(cells) <= min_required_idx:
                continue
            vals = [norm(c.get_text(" ", strip=True)) for c in cells]
            raw_product = vals[product_idx]
            product = clean_product(raw_product)
            if not product:
                continue
            low_p = product.lower()
            if low_p in {
                "jenis dokumen",
                "no.",
                "bilangan dokumen",
                "fi pemfailan (rm)",
                "jtt (rm)",
                "amaun deposit",
                "amaun deposit (rm)",
                "status dokumen",
            }:
                continue
            if "page size:" in low_p or "items in" in low_p:
                continue

            price = compute_effective_row_price(vals, price_idx, deposit_idx, jtt_idx)
            if price is None:
                continue

            person_from_prod = extract_person_from_product_text(raw_product)
            final_name = person_from_prod if person_from_prod else case_party_name

            row_id = tr.get("id", "") or f"{table.get('id', '')}_{len(results)}"
            key = (row_id, final_name.casefold(), product.casefold(), price)
            if key in seen:
                continue
            seen.add(key)
            results.append({
                "product": product,
                "price": price,
                "name": final_name,
            })

    return results


def get_case_details_container_html(page: Page) -> str:
    """Return the live CaseUserControl HTML after an EFS RowClick."""
    selectors = [
        "#ctl00_ctl00_Body_MyFilingBody_CaseUserControl",
        "#ctl00_ctl00_ctl00_ctl00_Body_MyFilingBody_CaseUserControlPanel",
        "[id$='CaseUserControl']",
        "[id*='CaseUserControlPanel']",
    ]
    for selector in selectors:
        try:
            loc = page.locator(selector)
            if loc.count():
                html = loc.first.inner_html(timeout=1500)
                if html and (
                    "Jenis Dokumen" in html
                    or "Fi Pemfailan" in html
                    or "Amaun Deposit" in html
                    or "radGridCaseDetails" in html
                    or "Butiran Pihak" in html
                ):
                    return html
        except Exception:
            continue
    return ""


def extract_live_document_rows(page: Page) -> list[dict[str, str]]:
    """Extract ALL document rows and their case party names from the live Case Details DOM."""
    js = r"""
    () => {
        const norm = (s) => (s || '').replace(/\s+/g, ' ').trim();
        const tables = Array.from(document.querySelectorAll('table'))
            .filter(t => {
                const r = t.getBoundingClientRect();
                return r.width > 0 && r.height > 0;
            });

        const candidates = [];

        for (const table of tables) {
            let headerRows = Array.from(table.children)
                .filter(x => x.tagName === 'THEAD')
                .flatMap(thead => Array.from(thead.children)
                    .filter(x => x.tagName === 'TR'));

            if (!headerRows.length) {
                headerRows = Array.from(table.children)
                    .filter(x => x.tagName === 'TR');
            }

            let header = null;
            for (let i = headerRows.length - 1; i >= 0; i--) {
                const cells = Array.from(headerRows[i].children)
                    .filter(x => x.tagName === 'TH' || x.tagName === 'TD')
                    .map(x => norm(x.innerText));
                const low = cells.map(x => x.toLowerCase());
                if (
                    low.some(x => x.includes('jenis dokumen')) &&
                    (low.some(x => x.includes('fi pemfailan')) || low.some(x => x.includes('deposit')))
                ) {
                    header = cells;
                    break;
                }
            }
            if (!header) continue;

            const lowHeader = header.map(x => x.toLowerCase());
            const productIdx = lowHeader.findIndex(x => x.includes('jenis dokumen'));
            const priceIdx = lowHeader.findIndex(x => x.includes('fi pemfailan'));
            const depositIdx = lowHeader.findIndex(x => x.includes('deposit'));
            const jttIdx = lowHeader.findIndex(x => x.includes('jtt'));
            if (productIdx < 0 || (priceIdx < 0 && depositIdx < 0)) continue;

            // Locate the master case row preceding this detail table
            let casePartyText = '';
            let detailTr = table.closest('tr');
            let curr = detailTr;
            while (curr) {
                curr = curr.previousElementSibling;
                if (curr && (curr.classList.contains('rgRow') || curr.classList.contains('rgAltRow'))) {
                    const tds = Array.from(curr.children).filter(x => x.tagName === 'TD');
                    if (tds.length >= 3) {
                        casePartyText = tds[2].innerText || '';
                    }
                    break;
                }
            }

            const minRequiredIdx = Math.max(productIdx, priceIdx, depositIdx);
            const rows = [];
            const tbodies = Array.from(table.children)
                .filter(x => x.tagName === 'TBODY');
            for (const tbody of tbodies) {
                for (const tr of Array.from(tbody.children)
                    .filter(x => x.tagName === 'TR')) {
                    const cells = Array.from(tr.children)
                        .filter(x => x.tagName === 'TD')
                        .map(x => norm(x.innerText));
                    if (cells.length > minRequiredIdx) {
                        rows.push({
                            id: tr.id || '',
                            cells: cells
                        });
                    }
                }
            }

            candidates.push({
                id: table.id || '',
                productIdx,
                priceIdx,
                depositIdx,
                jttIdx,
                rows,
                casePartyText,
            });
        }

        return candidates;
    }
    """

    try:
        candidates = page.evaluate(js)
    except Exception:
        candidates = []

    results: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for candidate in candidates or []:
        try:
            product_idx = int(candidate.get("productIdx", -1))
            price_idx = int(candidate.get("priceIdx", -1))
            deposit_idx = int(candidate.get("depositIdx", -1))
            jtt_idx = int(candidate.get("jttIdx", -1))
        except (TypeError, ValueError):
            continue

        case_party_text = str(candidate.get("casePartyText", ""))
        case_name = find_person_name(case_party_text)

        rows = candidate.get("rows") or []
        for row_item in rows:
            if isinstance(row_item, dict):
                row_id = row_item.get("id", "")
                vals = row_item.get("cells", [])
            elif isinstance(row_item, list):
                row_id = ""
                vals = row_item
            else:
                continue

            if not isinstance(vals, list):
                continue
            min_req = max(product_idx, price_idx, deposit_idx)
            if len(vals) <= min_req:
                continue

            raw_product = norm(str(vals[product_idx]))
            product = clean_product(raw_product)
            if not product:
                continue

            low = product.lower()
            if low in {
                "jenis dokumen",
                "no.",
                "bilangan dokumen",
                "fi pemfailan (rm)",
                "jtt (rm)",
                "amaun deposit",
                "amaun deposit (rm)",
                "status dokumen",
            }:
                continue
            if "page size:" in low or "items in" in low:
                continue

            str_vals = [norm(str(v)) for v in vals]
            price = compute_effective_row_price(str_vals, price_idx, deposit_idx, jtt_idx)
            if price is None:
                continue

            person_from_prod = extract_person_from_product_text(raw_product)
            final_name = person_from_prod if person_from_prod else case_name

            unique_row_id = row_id or f"{candidate.get('id', '')}_{len(results)}"
            key = (unique_row_id, final_name.casefold(), product.casefold(), price)
            if key in seen:
                continue
            seen.add(key)
            results.append({
                "product": product,
                "price": price,
                "name": final_name,
            })

    if not results:
        container_html = get_case_details_container_html(page)
        if container_html:
            results = extract_document_rows_from_case_html(container_html)

    return results


def extract_live_name(page: Page) -> str:
    html = get_case_details_container_html(page)
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    text = soup.get_text("\n", strip=True)
    return find_person_name(text)


def wait_for_live_case_details(page: Page, before_html: str) -> str:
    """Wait until the live page contains the document-grid headers."""
    try:
        page.wait_for_function(
            """({beforeHtml}) => {
                const text = document.body ? (document.body.innerText || '') : '';
                if (!text.includes('Jenis Dokumen') || (!text.includes('Fi Pemfailan') && !text.includes('Deposit'))) {
                    return false;
                }
                for (const table of document.querySelectorAll('table')) {
                    const r = table.getBoundingClientRect();
                    if (r.width <= 0 || r.height <= 0) continue;
                    const own = Array.from(table.children)
                        .filter(x => x.tagName === 'THEAD')
                        .flatMap(t => Array.from(t.children).filter(x => x.tagName === 'TR'))
                        .map(tr => (tr.innerText || ''))
                        .join(' ');
                    if (own.includes('Jenis Dokumen') && (own.includes('Fi Pemfailan') || own.includes('Deposit'))) {
                        return true;
                    }
                }
                const ids = [
                    'ctl00_ctl00_Body_MyFilingBody_CaseUserControl',
                    'ctl00_ctl00_ctl00_ctl00_Body_MyFilingBody_CaseUserControlPanel'
                ];
                for (const id of ids) {
                    const el = document.getElementById(id);
                    if (el && (el.innerHTML || '') !== beforeHtml) return true;
                }
                return false;
            }""",
            arg={"beforeHtml": before_html},
            timeout=15000,
        )
    except PlaywrightTimeoutError:
        pass
    return get_case_details_container_html(page)


def response_is_efs_ajax(response: Response) -> bool:
    try:
        headers = {k.lower(): v for k, v in response.request.headers.items()}
        return (
            response.request.method == "POST"
            and "/EFSWeb/eFiling/MyFiling.aspx" in response.url
            and headers.get("x-microsoftajax", "").lower().startswith("delta")
        )
    except Exception:
        return False


def click_filing_row(page: Page, index: int) -> dict[str, Any]:
    rows = page.locator(ROWS_SELECTOR)
    row = rows.nth(index)
    vals = [norm(x) for x in row.locator(":scope > td").all_inner_texts()]

    filing = {
        "filing_id": vals[0] if len(vals) > 0 else "",
        "court_full": vals[1] if len(vals) > 1 else "",
        "total_documents": vals[2] if len(vals) > 2 else "",
        "extraction_code": vals[3] if len(vals) > 3 else "",
        "total_amount": vals[4] if len(vals) > 4 else "",
        "status": vals[5] if len(vals) > 5 else "",
        "filing_date": extract_filing_date(vals[5]) if len(vals) > 5 else "",
        "row_index": str(index),
    }

    before_html = get_case_details_container_html(page)

    try:
        with page.expect_response(response_is_efs_ajax, timeout=12000):
            row.click(timeout=8000)
    except PlaywrightTimeoutError:
        pass

    wait_for_live_case_details(page, before_html)
    docs = extract_live_document_rows(page)

    if not docs:
        for delay in (150, 300, 500, 800):
            page.wait_for_timeout(delay)
            docs = extract_live_document_rows(page)
            if docs:
                break

    # Safety fallback: if a single-document filing (e.g., Deposit) still shows 0.00
    # while the master grid row has a positive payable amount, use the master amount.
    if len(docs) == 1 and docs[0].get("price") == "0.00":
        master_amt = parse_money_decimal(filing.get("total_amount", ""))
        if master_amt is not None and master_amt > Decimal("0.00"):
            docs[0]["price"] = f"{master_amt:.2f}"

    name = extract_live_name(page)
    court = court_code(filing["court_full"])
    filing_date = filing.get("filing_date", "")

    records = []
    previous_name = ""
    for i, doc in enumerate(docs):
        doc_name = clean_person_name(doc.get("name", "")) if doc.get("name") else ""
        if not doc_name:
            doc_name = name if i == 0 else previous_name

        # Every captured document row stores its filing DATE and full party NAME
        # so consecutive rows for the same person can be merged in Excel while
        # preserving underlying cell values in every row for filtering by PRODUCT/NAME.
        records.append({
            "DATE": filing_date,
            "NAME": doc_name,
            "PRODUCT": doc["product"],
            "COURT": court if i == 0 else "",
            "PRICE": doc["price"],
            "TOTAL": "",
        })
        previous_name = doc_name

    if docs:
        total = Decimal("0.00")
        for doc in docs:
            try:
                total += Decimal(doc["price"])
            except (InvalidOperation, TypeError, ValueError):
                pass
        records.append({
            "DATE": "",
            "NAME": "",
            "PRODUCT": "",
            "COURT": "",
            "PRICE": "",
            "TOTAL": f"{total:.2f}",
        })

    return {"filing": filing, "records": records}


def click_previous_page(page: Page) -> bool:
    current = current_page_number(page)
    if current <= 1:
        return False

    previous = page.locator(PREVIOUS_BUTTON_SELECTOR)
    if not previous.count():
        return False

    cls = previous.first.get_attribute("class") or ""
    if "rgDisabled" in cls:
        return False

    before_ids = page_rows(page)
    before_first = before_ids[0]["filing_id"] if before_ids else ""

    previous.first.click(timeout=10000)

    try:
        page.wait_for_function(
            """({beforePage, beforeFirst}) => {
                const id = "ctl00_ctl00_Body_MyFilingBody_radGridMyFiling";
                const g = window.$find ? window.$find(id) : null;
                if (g) {
                    const n = g.get_masterTableView().get_currentPageIndex() + 1;
                    if (n === beforePage - 1) return true;
                }
                const firstRow = document.querySelector(
                    "#ctl00_ctl00_Body_MyFilingBody_radGridMyFiling_ctl00 tbody tr.rgRow"
                );
                const first = firstRow ? (firstRow.innerText || "") : "";
                return beforeFirst && first && !first.includes(beforeFirst);
            }""",
            arg={"beforePage": current, "beforeFirst": before_first},
            timeout=15000,
        )
    except PlaywrightTimeoutError:
        page.wait_for_timeout(600)

    return current_page_number(page) == current - 1


def click_next_page(page: Page) -> bool:
    """Move forward (to an older page, page_no + 1) when seeking the start of a date range."""
    current = current_page_number(page)
    state = get_grid_client_state(page)
    page_count = int(state.get("pageCount", 0)) if state.get("ok") and state.get("pageCount") else 0
    if page_count and current >= page_count:
        return False

    nxt = page.locator(NEXT_BUTTON_SELECTOR)
    if not nxt.count():
        return False

    cls = nxt.first.get_attribute("class") or ""
    if "rgDisabled" in cls:
        return False

    before_ids = page_rows(page)
    before_first = before_ids[0]["filing_id"] if before_ids else ""

    nxt.first.click(timeout=10000)

    try:
        page.wait_for_function(
            """({beforePage, beforeFirst}) => {
                const id = "ctl00_ctl00_Body_MyFilingBody_radGridMyFiling";
                const g = window.$find ? window.$find(id) : null;
                if (g) {
                    const n = g.get_masterTableView().get_currentPageIndex() + 1;
                    if (n === beforePage + 1) return true;
                }
                const firstRow = document.querySelector(
                    "#ctl00_ctl00_Body_MyFilingBody_radGridMyFiling_ctl00 tbody tr.rgRow"
                );
                const first = firstRow ? (firstRow.innerText || "") : "";
                return beforeFirst && first && !first.includes(beforeFirst);
            }""",
            arg={"beforePage": current, "beforeFirst": before_first},
            timeout=15000,
        )
    except PlaywrightTimeoutError:
        page.wait_for_timeout(600)

    return current_page_number(page) == current + 1


def apply_date_grouping(records: list[dict[str, str]]) -> list[dict[str, str]]:
    """Preserve DATE on every captured document row without blanking out repeated dates."""
    return [dict(r) for r in records]


def cleanup_non_xlsx_files() -> None:
    """Ensure only final .xlsx files remain in the efs_output folder."""
    if not OUTPUT_DIR.exists():
        return
    for item in OUTPUT_DIR.iterdir():
        if not item.is_file():
            continue
        if item.suffix.lower() != ".xlsx" or item.name == "efs_reverse_checkpoint.xlsx":
            try:
                item.unlink()
            except Exception:
                pass


def export_to_excel(records: list[dict[str, str]], path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EFS Filing List"
    ws.views.sheetView[0].showGridLines = True

    # Page setup for Landscape A4 printing
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.freeze_panes = "A2"

    # Page margins (in inches) for clean landscape fit
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45

    # Typography & Color Palette
    font_name = "Calibri"
    header_font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")

    data_font = Font(name=font_name, size=10.5, color="111827")
    total_font = Font(name=font_name, size=11, bold=True, color="1B365D")
    total_fill = PatternFill(start_color="EBF1F8", end_color="EBF1F8", fill_type="solid")

    thin_gray = Side(style="thin", color="CBD5E1")
    double_bottom = Side(style="double", color="1B365D")
    top_border = Side(style="thin", color="1B365D")

    data_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    total_border = Border(left=thin_gray, right=thin_gray, top=top_border, bottom=double_bottom)
    header_border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # Header Row
    ws.append(OUTPUT_HEADERS)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(OUTPUT_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = align_center if col_idx == 1 else (align_right if col_idx in (5, 6) else align_left)

    # Ensure every document row within a filing carries its party NAME
    grouped = apply_date_grouping(records)
    normalized_rows: list[dict[str, str]] = []
    active_name = ""
    for r in grouped:
        item = dict(r)
        is_tot = bool(item.get("TOTAL")) and not bool(item.get("PRODUCT"))
        if is_tot:
            active_name = ""
        else:
            if item.get("NAME"):
                active_name = item["NAME"]
            elif active_name:
                item["NAME"] = active_name
        normalized_rows.append(item)

    align_name_merged = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data Rows
    for row_idx, r in enumerate(normalized_rows, start=2):
        is_total_row = bool(r.get("TOTAL")) and not bool(r.get("PRODUCT"))

        price_raw = r.get("PRICE", "")
        total_raw = r.get("TOTAL", "")
        try:
            price_val = float(price_raw) if price_raw else ""
        except (ValueError, TypeError):
            price_val = price_raw

        try:
            total_val = float(total_raw) if total_raw else ""
        except (ValueError, TypeError):
            total_val = total_raw

        name_txt = r.get("NAME", "")
        prod_txt = r.get("PRODUCT", "")
        court_txt = r.get("COURT", "")

        row_values = [
            r.get("DATE", ""),
            name_txt,
            prod_txt,
            court_txt,
            price_val,
            total_val,
        ]

        # Adjust row height dynamically if text wraps onto multiple lines
        if is_total_row:
            ws.row_dimensions[row_idx].height = 23
        elif len(name_txt) > 36 or len(prod_txt) > 34 or len(court_txt) > 52:
            ws.row_dimensions[row_idx].height = 32
        else:
            ws.row_dimensions[row_idx].height = 22

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if is_total_row:
                cell.font = total_font
                cell.fill = total_fill
                cell.border = total_border
            else:
                cell.font = data_font
                cell.border = data_border

            if col_idx == 1:
                cell.alignment = align_center
            elif col_idx == 2:
                cell.alignment = align_name_merged
            elif col_idx in (5, 6):
                cell.alignment = align_right
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = align_left

    # Merge consecutive NAME cells (Column B) for the same person within a filing.
    # Using ws.merged_cells.add() merges the visual block while keeping the person
    # name stored in every underlying cell so Excel/WPS filtering by PRODUCT or NAME
    # always returns the full list of person names and products!
    run_start: int | None = None
    run_name = ""
    for row_idx, r in enumerate(normalized_rows, start=2):
        is_total_row = bool(r.get("TOTAL")) and not bool(r.get("PRODUCT"))
        cur_name = r.get("NAME", "").strip() if not is_total_row else ""

        if cur_name and cur_name == run_name and run_start is not None:
            continue
        else:
            if run_start is not None and (row_idx - 1) > run_start:
                ws.merged_cells.add(f"B{run_start}:B{row_idx - 1}")
            if cur_name:
                run_start = row_idx
                run_name = cur_name
            else:
                run_start = None
                run_name = ""

    if run_start is not None and (len(normalized_rows) + 1) > run_start:
        ws.merged_cells.add(f"B{run_start}:B{len(normalized_rows) + 1}")

    # Column Widths optimized for A4 Landscape printing
    fixed_widths = {
        "A": 14,  # DATE
        "B": 38,  # NAME
        "C": 36,  # PRODUCT
        "D": 54,  # COURT
        "E": 15,  # PRICE
        "F": 15,  # TOTAL
    }
    for col_letter, width in fixed_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(str(path))
    return path


def save_checkpoint(records: list[dict[str, str]], page_no: int) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = OUTPUT_DIR / "efs_reverse_checkpoint.xlsx"
    export_to_excel(records, xlsx_path)
    return xlsx_path


def save_final(records: list[dict[str, str]], suffix_tag: str = "") -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = time.strftime("%Y%m%d_%H%M%S")
    tag_part = f"_{suffix_tag}" if suffix_tag else ""
    xlsx_path = OUTPUT_DIR / f"efs_result_reverse{tag_part}_{ts}.xlsx"
    export_to_excel(records, xlsx_path)
    cleanup_non_xlsx_files()
    return xlsx_path


def save_test(records: list[dict[str, str]]) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = OUTPUT_DIR / "efs_test_result.xlsx"
    export_to_excel(records, xlsx_path)
    cleanup_non_xlsx_files()
    return xlsx_path


def open_output_folder() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(OUTPUT_DIR.resolve()))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(OUTPUT_DIR.resolve())])
        else:
            subprocess.Popen(["xdg-open", str(OUTPUT_DIR.resolve())])
    except Exception:
        pass


# ==============================================================================
# DEDICATED PLAYWRIGHT WORKER & MODERN DESKTOP UI (FINAL PRODUCT EDITION)
# ==============================================================================


class PlaywrightWorker(threading.Thread):
    """Runs all Playwright browser operations on a single dedicated thread."""

    def __init__(self, cmd_q: queue.Queue[tuple[str, Any]], ui_q: queue.Queue[dict[str, Any]]) -> None:
        super().__init__(daemon=True)
        self.cmd_q = cmd_q
        self.ui_q = ui_q
        self.stop_event = threading.Event()
        self.playwright = None
        self.context = None
        self.page: Page | None = None
        self.temp_profile_dir: Path | None = None

    def emit(self, event_type: str, **kwargs: Any) -> None:
        payload = {"type": event_type}
        payload.update(kwargs)
        self.ui_q.put(payload)

    def ensure_browser(self) -> Page | None:
        try:
            if self.page is not None and not self.page.is_closed():
                return self.page
        except Exception:
            self.page = None

        try:
            if self.playwright is None:
                self.playwright = sync_playwright().start()
            self.emit("STATUS", text="Launching Microsoft Edge...", state="busy")
            self.context, self.page, self.temp_profile_dir = launch_edge_resilient(
                self.playwright,
                log_fn=lambda msg: self.emit("LOG", message=msg),
            )
            self.page.goto(URL, wait_until="domcontentloaded", timeout=30000)
            self.emit(
                "LOG",
                message="Browser opened. Please log in to EFS, open My Filing, set Page Size = 10, and navigate to your starting page.",
            )
            self.emit("BROWSER_READY", text="Browser Connected — Ready to Scrape")
            return self.page
        except Exception as exc:
            self.emit("ERROR", message=f"Failed to launch Microsoft Edge: {exc}")
            return None

    def run(self) -> None:
        while True:
            cmd, arg = self.cmd_q.get()
            if cmd == "QUIT":
                self._cleanup()
                break
            elif cmd == "LAUNCH":
                self.ensure_browser()
            elif cmd in {"TEST_LOWEST", "SCRAPE_PAGE", "SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"}:
                self.stop_event.clear()
                self._execute_scrape(cmd, arg or {})

    def _seek_start_date_page(self, page: Page, start_dt: date) -> None:
        """If the current page only contains filings newer than start_dt, advance pages
        (using Next Page) until we reach the earliest page containing start_dt."""
        while not self.stop_event.is_set():
            page_no = current_page_number(page)
            rows = page_rows(page)
            if not rows:
                break
            row_dates = [parse_user_date(r["filing_date"]) for r in rows]
            valid_dates = [d for d in row_dates if d is not None]
            if not valid_dates:
                break

            oldest_on_page = min(valid_dates)
            if oldest_on_page >= start_dt:
                self.emit(
                    "LOG",
                    message=(
                        f"Auto-seeking to From Date ({start_dt.strftime('%d/%m/%Y')}): "
                        f"Page {page_no} oldest filing is {oldest_on_page.strftime('%d/%m/%Y')} — advancing to Page {page_no + 1}..."
                    ),
                )
                self.emit(
                    "PROGRESS",
                    page=page_no,
                    step=1,
                    total_steps=1,
                    filing_id=f"Seeking {start_dt.strftime('%d/%m/%Y')}...",
                )
                if not click_next_page(page):
                    break
            else:
                break

    def _execute_scrape(self, mode: str, options: dict[str, Any]) -> None:
        page = self.ensure_browser()
        if page is None:
            return

        start_dt: date | None = options.get("start_date")
        end_dt: date | None = options.get("end_date")
        auto_seek: bool = bool(options.get("auto_seek", True))

        self.emit("SCRAPE_STARTED", mode=mode)
        try:
            page.wait_for_selector(ROWS_SELECTOR, timeout=15000)
        except PlaywrightTimeoutError:
            self.emit(
                "ERROR",
                message="Filing grid not detected. Please make sure 'My Filing' is open in Microsoft Edge.",
            )
            self.emit("SCRAPE_FINISHED")
            return

        page_no = current_page_number(page)
        page_size = current_page_size(page)
        rows = page_rows(page)

        if mode in {"SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"}:
            if page_size != "10":
                self.emit(
                    "ERROR",
                    message=f"Page Size is currently '{page_size or 'unknown'}'. Please set EFS Page Size = 10 before starting.",
                )
                self.emit("SCRAPE_FINISHED")
                return
            if mode == "SCRAPE_REVERSE" and start_dt is None and end_dt is None and page_no == 1:
                self.emit(
                    "ERROR",
                    message="You are currently on Page 1. For Reverse Scrape without a Date Range, please navigate to your older starting page (e.g. Page 293) first, or enter a Date Range.",
                )
                self.emit("SCRAPE_FINISHED")
                return

        if start_dt or end_dt:
            s_label = start_dt.strftime("%d/%m/%Y") if start_dt else "Earliest"
            e_label = end_dt.strftime("%d/%m/%Y") if end_dt else "Latest"
            self.emit("LOG", message=f"Date Range Filter Active: {s_label}  ->  {e_label}")

        if mode == "TEST_LOWEST":
            if not rows:
                self.emit("ERROR", message="No filing rows found on the current page.")
                self.emit("SCRAPE_FINISHED")
                return

            test_idx = len(rows) - 1
            filing_id = rows[test_idx]["filing_id"]
            self.emit("LOG", message=f"Testing lowest filing on Page {page_no} (Row {test_idx + 1}: {filing_id})...")
            self.emit("PROGRESS", page=page_no, step=1, total_steps=1, filing_id=filing_id)

            started = time.monotonic()
            result = click_filing_row(page, test_idx)
            elapsed = time.monotonic() - started
            records = apply_date_grouping(result["records"])

            self.emit("RECORDS_BATCH", records=records, filing_id=filing_id, elapsed=elapsed)
            xlsx_path = save_test(records)
            self.emit(
                "DONE",
                xlsx_path=str(xlsx_path.resolve()),
                total_rows=len(records),
                mode=mode,
            )
            self.emit("SCRAPE_FINISHED")
            return

        # If a From Date (start_dt) is specified and auto_seek is enabled (in multi-page modes),
        # automatically advance forward if the current page is newer than start_dt.
        if mode in {"SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"} and start_dt is not None and auto_seek:
            self._seek_start_date_page(page, start_dt)

        all_records: list[dict[str, str]] = []
        processed_ids: set[str] = set()
        reached_past_end_date = False

        while not self.stop_event.is_set() and not reached_past_end_date:
            page_no = current_page_number(page)
            rows = page_rows(page)
            total_on_page = len(rows)
            if not rows:
                break

            row_dates = [parse_user_date(r["filing_date"]) for r in rows]
            valid_dates = [d for d in row_dates if d is not None]

            # Page-level fast skip: if every row on this page is older than start_dt
            if (
                mode in {"SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"}
                and start_dt is not None
                and valid_dates
                and max(valid_dates) < start_dt
            ):
                self.emit(
                    "LOG",
                    message=(
                        f"Page {page_no}: All filings ({max(valid_dates).strftime('%d/%m/%Y')}) are older than "
                        f"From Date ({start_dt.strftime('%d/%m/%Y')}) — skipping page..."
                    ),
                )
                if page_no <= 1 or not click_previous_page(page):
                    break
                continue

            # Page-level fast stop: if every row on this page is already newer than end_dt
            if (
                mode in {"SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"}
                and end_dt is not None
                and valid_dates
                and min(valid_dates) > end_dt
            ):
                self.emit(
                    "LOG",
                    message=(
                        f"Page {page_no}: Oldest filing ({min(valid_dates).strftime('%d/%m/%Y')}) is newer than "
                        f"To Date ({end_dt.strftime('%d/%m/%Y')}). Date range scan complete!"
                    ),
                )
                break

            self.emit("LOG", message=f"Processing Page {page_no} ({total_on_page} filings, lowest -> highest)...")

            row_order = list(range(total_on_page - 1, -1, -1))
            for step, i in enumerate(row_order, start=1):
                if self.stop_event.is_set():
                    self.emit("LOG", message="Stop requested by user. Saving captured records...")
                    break

                filing_id = rows[i]["filing_id"]
                row_date_str = rows[i].get("filing_date", "")
                row_dt = parse_user_date(row_date_str)

                if filing_id in processed_ids:
                    continue

                # Date Range Checks (without clicking out-of-range rows!)
                if start_dt is not None and row_dt is not None and row_dt < start_dt:
                    self.emit(
                        "LOG",
                        message=f"  [Skip] Row {i + 1} ({filing_id}, {row_date_str}) is older than {start_dt.strftime('%d/%m/%Y')}",
                    )
                    continue

                if end_dt is not None and row_dt is not None and row_dt > end_dt:
                    self.emit(
                        "LOG",
                        message=f"  [Stop] Row {i + 1} ({filing_id}, {row_date_str}) is newer than To Date ({end_dt.strftime('%d/%m/%Y')}).",
                    )
                    if mode in {"SCRAPE_REVERSE", "SCRAPE_DATE_RANGE"}:
                        reached_past_end_date = True
                        break
                    continue

                self.emit(
                    "PROGRESS",
                    page=page_no,
                    step=step,
                    total_steps=total_on_page,
                    filing_id=f"{filing_id} ({row_date_str})",
                )
                try:
                    started = time.monotonic()
                    result = click_filing_row(page, i)
                    elapsed = time.monotonic() - started
                    processed_ids.add(filing_id)
                    batch = result["records"]
                    all_records.extend(batch)
                    self.emit(
                        "RECORDS_BATCH",
                        records=batch,
                        filing_id=filing_id,
                        elapsed=elapsed,
                        row_num=i + 1,
                        step=step,
                        total_steps=total_on_page,
                    )
                except Exception as exc:
                    self.emit("LOG", message=f"Row {i + 1} ({filing_id}) failed: {exc}")

            if mode == "SCRAPE_PAGE" or self.stop_event.is_set() or reached_past_end_date:
                break

            if all_records:
                save_checkpoint(all_records, page_no)

            if page_no <= 1:
                break
            self.emit("LOG", message=f"Moving from Page {page_no} to Page {page_no - 1}...")
            if not click_previous_page(page):
                self.emit("LOG", message="Could not move to previous page. Stopping safely.")
                break

        if all_records:
            suffix = ""
            if start_dt or end_dt:
                s_tag = start_dt.strftime("%Y%m%d") if start_dt else "start"
                e_tag = end_dt.strftime("%Y%m%d") if end_dt else "end"
                suffix = f"{s_tag}_to_{e_tag}"
            xlsx_path = save_final(all_records, suffix_tag=suffix)
            self.emit(
                "DONE",
                xlsx_path=str(xlsx_path.resolve()),
                total_rows=len(all_records),
                mode=mode,
            )
        else:
            cleanup_non_xlsx_files()
            self.emit("LOG", message="No matching document rows were captured in the selected range.")
        self.emit("SCRAPE_FINISHED")

    def _cleanup(self) -> None:
        try:
            if self.context is not None:
                self.context.close()
        except Exception:
            pass
        try:
            if self.playwright is not None:
                self.playwright.stop()
        except Exception:
            pass
        if self.temp_profile_dir and self.temp_profile_dir.exists():
            shutil.rmtree(self.temp_profile_dir, ignore_errors=True)


def launch_gui() -> bool:
    """Launch the polished Tkinter Desktop UI. Returns False if Tkinter is unavailable."""
    try:
        import tkinter as tk
        from tkinter import messagebox, ttk
    except Exception:
        return False

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    cmd_q: queue.Queue[tuple[str, Any]] = queue.Queue()
    ui_q: queue.Queue[dict[str, Any]] = queue.Queue()
    worker = PlaywrightWorker(cmd_q, ui_q)
    worker.start()

    root = tk.Tk()
    root.title(APP_TITLE)
    root.geometry("1180x810")
    root.minsize(1000, 680)
    root.configure(bg="#F4F6F9")

    # Configure ttk Theme & Styles
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except Exception:
        pass

    style.configure("Treeview", font=("Segoe UI", 9), rowheight=25, background="#FFFFFF", fieldbackground="#FFFFFF")
    style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"), background="#1B365D", foreground="#FFFFFF", padding=6)
    style.map("Treeview.Heading", background=[("active", "#24477A")])

    # ── Top Header Banner ──
    header_frame = tk.Frame(root, bg="#1B365D", height=68)
    header_frame.pack(fill="x", side="top")
    header_frame.pack_propagate(False)

    title_box = tk.Frame(header_frame, bg="#1B365D")
    title_box.pack(side="left", padx=20, pady=10)
    tk.Label(
        title_box,
        text="MALAYSIA COURT E-FILING (EFS) SCRAPER",
        font=("Segoe UI", 15, "bold"),
        fg="#FFFFFF",
        bg="#1B365D",
    ).pack(anchor="w")
    tk.Label(
        title_box,
        text="Automated Case, Party, Deposit & Date-Range Extraction System  •  Landscape A4 Excel Export",
        font=("Segoe UI", 9),
        fg="#CBD5E1",
        bg="#1B365D",
    ).pack(anchor="w")

    status_pill = tk.Label(
        header_frame,
        text="● Browser: Not Launched",
        font=("Segoe UI", 9, "bold"),
        fg="#FCD34D",
        bg="#24477A",
        padx=14,
        pady=6,
    )
    status_pill.pack(side="right", padx=20, pady=16)

    # ── Control Toolbar ──
    toolbar = tk.Frame(root, bg="#FFFFFF", bd=1, relief="solid")
    toolbar.pack(fill="x", padx=16, pady=(12, 6))

    btn_font = ("Segoe UI", 9, "bold")

    def make_btn(parent: tk.Widget, text: str, bg_color: str, fg_color: str, command: Any, width: int = 20) -> tk.Button:
        return tk.Button(
            parent,
            text=text,
            font=btn_font,
            bg=bg_color,
            fg=fg_color,
            activebackground=bg_color,
            activeforeground=fg_color,
            relief="flat",
            cursor="hand2",
            padx=10,
            pady=7,
            width=width,
            command=command,
        )

    # Date Range Variables
    from_date_var = tk.StringVar(value="")
    to_date_var = tk.StringVar(value="")
    auto_seek_var = tk.BooleanVar(value=True)

    def get_validated_date_options(require_dates: bool = False) -> dict[str, Any] | None:
        s_raw = from_date_var.get().strip()
        e_raw = to_date_var.get().strip()
        if require_dates and not s_raw and not e_raw:
            messagebox.showwarning(
                "Date Range Required",
                "Please enter a 'From Date' and/or 'To Date' (DD/MM/YYYY, e.g. 05/10/2026) before clicking 'Scan Date Range'.",
            )
            return None
        s_dt, e_dt, err = normalize_date_range(s_raw, e_raw)
        if err:
            messagebox.showwarning("Invalid Date Format", err)
            return None
        if s_dt and e_dt:
            # Update entry boxes with normalized DD/MM/YYYY order
            from_date_var.set(s_dt.strftime("%d/%m/%Y"))
            to_date_var.set(e_dt.strftime("%d/%m/%Y"))
        return {
            "start_date": s_dt,
            "end_date": e_dt,
            "auto_seek": bool(auto_seek_var.get()),
        }

    btn_row = tk.Frame(toolbar, bg="#FFFFFF")
    btn_row.pack(fill="x", padx=12, pady=(10, 6))

    def on_launch_browser() -> None:
        cmd_q.put(("LAUNCH", None))

    def on_start_reverse() -> None:
        opts = get_validated_date_options(require_dates=False)
        if opts is None:
            return
        cmd_q.put(("SCRAPE_REVERSE", opts))

    def on_scan_date_range() -> None:
        opts = get_validated_date_options(require_dates=True)
        if opts is None:
            return
        cmd_q.put(("SCRAPE_DATE_RANGE", opts))

    def on_scrape_page() -> None:
        opts = get_validated_date_options(require_dates=False)
        if opts is None:
            return
        cmd_q.put(("SCRAPE_PAGE", opts))

    def on_test_lowest() -> None:
        cmd_q.put(("TEST_LOWEST", None))

    def on_stop() -> None:
        worker.stop_event.set()
        btn_stop.configure(state="disabled")

    def on_clear_table() -> None:
        for item in tree.get_children():
            tree.delete(item)
        state_vars["filings"].set("0")
        state_vars["rows"].set("0")
        state_vars["total_rm"].set("RM 0.00")
        state_vars["accum_rm"] = Decimal("0.00")
        state_vars["filings_count"] = 0
        state_vars["rows_count"] = 0

    btn_launch = make_btn(btn_row, "1. Launch Edge Browser", "#1B365D", "#FFFFFF", on_launch_browser, width=19)
    btn_launch.pack(side="left", padx=(0, 6))

    btn_reverse = make_btn(btn_row, "▶ Reverse Scrape (To Page 1)", "#15803D", "#FFFFFF", on_start_reverse, width=23)
    btn_reverse.pack(side="left", padx=4)

    btn_date_scan = make_btn(btn_row, "📅 Scan Date Range", "#7C3AED", "#FFFFFF", on_scan_date_range, width=18)
    btn_date_scan.pack(side="left", padx=4)

    btn_page = make_btn(btn_row, "📄 Scrape Current Page", "#1E56A0", "#FFFFFF", on_scrape_page, width=19)
    btn_page.pack(side="left", padx=4)

    btn_test = make_btn(btn_row, "🔍 Test Lowest", "#475569", "#FFFFFF", on_test_lowest, width=13)
    btn_test.pack(side="left", padx=4)

    btn_stop = make_btn(btn_row, "⏹ Stop & Save", "#DC2626", "#FFFFFF", on_stop, width=13)
    btn_stop.configure(state="disabled")
    btn_stop.pack(side="left", padx=4)

    btn_folder = make_btn(btn_row, "📂 Open efs_output", "#0F766E", "#FFFFFF", open_output_folder, width=16)
    btn_folder.pack(side="right", padx=(4, 0))

    # ── Date Range Filter Sub-Bar ──
    date_bar = tk.Frame(toolbar, bg="#F8FAFC", bd=1, relief="solid")
    date_bar.pack(fill="x", padx=12, pady=(2, 10))

    tk.Label(
        date_bar,
        text="📅 DATE RANGE FILTER (DD/MM/YYYY):",
        font=("Segoe UI", 9, "bold"),
        fg="#1B365D",
        bg="#F8FAFC",
    ).pack(side="left", padx=(10, 8), pady=7)

    tk.Label(date_bar, text="From Date:", font=("Segoe UI", 9), fg="#334155", bg="#F8FAFC").pack(side="left", padx=(4, 4))
    from_entry = tk.Entry(
        date_bar,
        textvariable=from_date_var,
        font=("Segoe UI", 9, "bold"),
        width=13,
        bd=1,
        relief="solid",
        justify="center",
    )
    from_entry.pack(side="left", padx=(0, 10), pady=6, ipady=3)

    tk.Label(date_bar, text="To Date:", font=("Segoe UI", 9), fg="#334155", bg="#F8FAFC").pack(side="left", padx=(2, 4))
    to_entry = tk.Entry(
        date_bar,
        textvariable=to_date_var,
        font=("Segoe UI", 9, "bold"),
        width=13,
        bd=1,
        relief="solid",
        justify="center",
    )
    to_entry.pack(side="left", padx=(0, 10), pady=6, ipady=3)

    def set_today_dates() -> None:
        today_str = datetime.now().strftime("%d/%m/%Y")
        from_date_var.set(today_str)
        to_date_var.set(today_str)

    def clear_dates() -> None:
        from_date_var.set("")
        to_date_var.set("")

    tk.Button(
        date_bar,
        text="Today",
        font=("Segoe UI", 8, "bold"),
        bg="#E2E8F0",
        fg="#1E293B",
        relief="flat",
        cursor="hand2",
        padx=8,
        pady=3,
        command=set_today_dates,
    ).pack(side="left", padx=3)

    tk.Button(
        date_bar,
        text="Clear Dates",
        font=("Segoe UI", 8, "bold"),
        bg="#E2E8F0",
        fg="#1E293B",
        relief="flat",
        cursor="hand2",
        padx=8,
        pady=3,
        command=clear_dates,
    ).pack(side="left", padx=3)

    tk.Checkbutton(
        date_bar,
        text="Auto-seek forward to From Date if current page is newer",
        variable=auto_seek_var,
        font=("Segoe UI", 8),
        fg="#475569",
        bg="#F8FAFC",
        activebackground="#F8FAFC",
    ).pack(side="left", padx=12)

    tk.Button(
        date_bar,
        text="Clear Table Preview",
        font=("Segoe UI", 8, "bold"),
        bg="#E2E8F0",
        fg="#1E293B",
        relief="flat",
        cursor="hand2",
        padx=10,
        pady=3,
        command=on_clear_table,
    ).pack(side="right", padx=10)

    # ── KPI Summary Cards ──
    kpi_frame = tk.Frame(root, bg="#F4F6F9")
    kpi_frame.pack(fill="x", padx=16, pady=(0, 8))

    state_vars: dict[str, Any] = {
        "page": tk.StringVar(value="—"),
        "filings": tk.StringVar(value="0"),
        "rows": tk.StringVar(value="0"),
        "total_rm": tk.StringVar(value="RM 0.00"),
        "status_text": tk.StringVar(
            value="Step 1: Click '1. Launch Edge Browser', log in to EFS (Page Size = 10), optionally set a Date Range, and click 'Scan Date Range' or 'Reverse Scrape'."
        ),
        "accum_rm": Decimal("0.00"),
        "filings_count": 0,
        "rows_count": 0,
    }

    kpi_specs = [
        ("CURRENT EFS PAGE", state_vars["page"], "#1B365D"),
        ("FILINGS PROCESSED", state_vars["filings"], "#1E56A0"),
        ("DOCUMENT ROWS CAPTURED", state_vars["rows"], "#0F766E"),
        ("ACCUMULATED TOTAL (RM)", state_vars["total_rm"], "#15803D"),
    ]

    for idx, (label_text, var, accent_color) in enumerate(kpi_specs):
        card = tk.Frame(kpi_frame, bg="#FFFFFF", bd=1, relief="solid")
        card.pack(side="left", fill="both", expand=True, padx=(0 if idx == 0 else 6, 0 if idx == len(kpi_specs) - 1 else 6))
        tk.Label(card, text=label_text, font=("Segoe UI", 8, "bold"), fg="#64748B", bg="#FFFFFF").pack(anchor="w", padx=14, pady=(8, 2))
        tk.Label(card, textvariable=var, font=("Segoe UI", 16, "bold"), fg=accent_color, bg="#FFFFFF").pack(anchor="w", padx=14, pady=(0, 8))

    # ── Progress & Status Bar ──
    prog_frame = tk.Frame(root, bg="#FFFFFF", bd=1, relief="solid")
    prog_frame.pack(fill="x", padx=16, pady=(0, 8))

    tk.Label(
        prog_frame,
        textvariable=state_vars["status_text"],
        font=("Segoe UI", 9, "bold"),
        fg="#1E293B",
        bg="#FFFFFF",
        anchor="w",
    ).pack(fill="x", padx=12, pady=(6, 4))

    progress_bar = ttk.Progressbar(prog_frame, orient="horizontal", mode="determinate", maximum=100)
    progress_bar.pack(fill="x", padx=12, pady=(0, 8))

    # ── Live Captured Rows Table ──
    table_container = tk.Frame(root, bg="#FFFFFF", bd=1, relief="solid")
    table_container.pack(fill="both", expand=True, padx=16, pady=(0, 8))

    cols = ("NO", "DATE", "NAME", "PRODUCT", "COURT", "PRICE", "TOTAL")
    tree = ttk.Treeview(table_container, columns=cols, show="headings", selectmode="browse")

    col_configs = [
        ("NO", "#", 45, "center"),
        ("DATE", "DATE", 95, "center"),
        ("NAME", "NAME", 240, "w"),
        ("PRODUCT", "PRODUCT", 220, "w"),
        ("COURT", "COURT", 320, "w"),
        ("PRICE", "PRICE (RM)", 100, "e"),
        ("TOTAL", "TOTAL (RM)", 100, "e"),
    ]
    for col_id, heading, width, anchor in col_configs:
        tree.heading(col_id, text=heading)
        tree.column(col_id, width=width, anchor=anchor, stretch=(col_id in {"NAME", "PRODUCT", "COURT"}))

    tree.tag_configure("odd", background="#FFFFFF")
    tree.tag_configure("even", background="#F8FAFC")
    tree.tag_configure("total_row", background="#EBF1F8", foreground="#1B365D", font=("Segoe UI", 9, "bold"))

    vsb = ttk.Scrollbar(table_container, orient="vertical", command=tree.yview)
    hsb = ttk.Scrollbar(table_container, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

    tree.grid(row=0, column=0, sticky="nsew")
    vsb.grid(row=0, column=1, sticky="ns")
    hsb.grid(row=1, column=0, sticky="ew")
    table_container.grid_rowconfigure(0, weight=1)
    table_container.grid_columnconfigure(0, weight=1)

    # ── Bottom Activity Log ──
    log_frame = tk.Frame(root, bg="#FFFFFF", bd=1, relief="solid", height=110)
    log_frame.pack(fill="x", padx=16, pady=(0, 12))
    log_frame.pack_propagate(False)

    log_text = tk.Text(
        log_frame,
        font=("Consolas", 9),
        bg="#0F172A",
        fg="#E2E8F0",
        relief="flat",
        padx=10,
        pady=6,
        state="disabled",
    )
    log_scroll = ttk.Scrollbar(log_frame, orient="vertical", command=log_text.yview)
    log_text.configure(yscrollcommand=log_scroll.set)
    log_text.pack(side="left", fill="both", expand=True)
    log_scroll.pack(side="right", fill="y")

    def append_log(msg: str) -> None:
        ts = datetime.now().strftime("%H:%M:%S")
        line = f"[{ts}] {msg}\n"
        print(line, end="")
        log_text.configure(state="normal")
        log_text.insert("end", line)
        log_text.see("end")
        log_text.configure(state="disabled")

    def set_buttons_scraping(is_scraping: bool) -> None:
        state_normal = "disabled" if is_scraping else "normal"
        btn_launch.configure(state=state_normal)
        btn_reverse.configure(state=state_normal)
        btn_date_scan.configure(state=state_normal)
        btn_page.configure(state=state_normal)
        btn_test.configure(state=state_normal)
        btn_stop.configure(state="normal" if is_scraping else "disabled")

    def poll_ui_queue() -> None:
        try:
            while True:
                msg = ui_q.get_nowait()
                mtype = msg.get("type", "")

                if mtype == "LOG":
                    append_log(msg.get("message", ""))

                elif mtype == "BROWSER_READY":
                    status_pill.configure(text="● Browser: Connected", fg="#4ADE80")
                    state_vars["status_text"].set(
                        "Browser connected. Log in to EFS (Page Size = 10), set an optional Date Range above, and click 'Scan Date Range' or 'Reverse Scrape'."
                    )

                elif mtype == "SCRAPE_STARTED":
                    set_buttons_scraping(True)
                    status_pill.configure(text="● Scraping Active...", fg="#38BDF8")
                    progress_bar["value"] = 0

                elif mtype == "PROGRESS":
                    p_no = msg.get("page", 0)
                    step = msg.get("step", 1)
                    total_steps = max(msg.get("total_steps", 1), 1)
                    fid = msg.get("filing_id", "")
                    state_vars["page"].set(f"Page {p_no}")
                    pct = int((step - 1) / total_steps * 100)
                    progress_bar["value"] = pct
                    state_vars["status_text"].set(
                        f"Page {p_no}  •  Processing Filing {step}/{total_steps} ({fid})  [Lowest → Highest]"
                    )

                elif mtype == "RECORDS_BATCH":
                    batch = msg.get("records", [])
                    fid = msg.get("filing_id", "")
                    elapsed = msg.get("elapsed", 0.0)
                    step = msg.get("step", 1)
                    total_steps = max(msg.get("total_steps", 1), 1)
                    progress_bar["value"] = int(step / total_steps * 100)

                    state_vars["filings_count"] += 1
                    state_vars["filings"].set(str(state_vars["filings_count"]))

                    doc_count = 0
                    for r in batch:
                        is_tot = bool(r.get("TOTAL")) and not bool(r.get("PRODUCT"))
                        if is_tot:
                            tot_dec = parse_money_decimal(r.get("TOTAL", "")) or Decimal("0.00")
                            state_vars["accum_rm"] += tot_dec
                            state_vars["total_rm"].set(f"RM {state_vars['accum_rm']:,.2f}")
                            tag = "total_row"
                            row_no_str = ""
                        else:
                            doc_count += 1
                            state_vars["rows_count"] += 1
                            tag = "even" if state_vars["rows_count"] % 2 == 0 else "odd"
                            row_no_str = str(state_vars["rows_count"])

                        item_id = tree.insert(
                            "",
                            "end",
                            values=(
                                row_no_str,
                                r.get("DATE", ""),
                                r.get("NAME", ""),
                                r.get("PRODUCT", ""),
                                r.get("COURT", ""),
                                r.get("PRICE", ""),
                                r.get("TOTAL", ""),
                            ),
                            tags=(tag,),
                        )
                        tree.see(item_id)

                    state_vars["rows"].set(str(state_vars["rows_count"]))
                    append_log(f"Captured {fid}: {doc_count} document row(s) in {elapsed:.2f}s")

                elif mtype == "DONE":
                    xlsx_path = msg.get("xlsx_path", "")
                    total_rows = msg.get("total_rows", 0)
                    progress_bar["value"] = 100
                    state_vars["status_text"].set(f"Completed! Saved Excel file: {xlsx_path}")
                    append_log(f"SUCCESS — Saved styled Excel (.xlsx) with {total_rows} rows: {xlsx_path}")
                    messagebox.showinfo(
                        "Export Complete",
                        f"Excel file generated successfully!\n\nFile:\n{xlsx_path}\n\nTotal rows written: {total_rows}",
                    )

                elif mtype == "ERROR":
                    err_msg = msg.get("message", "Unknown error")
                    append_log(f"ERROR: {err_msg}")
                    state_vars["status_text"].set(f"Error: {err_msg}")
                    messagebox.showwarning("EFS Scraper Notice", err_msg)

                elif mtype == "SCRAPE_FINISHED":
                    set_buttons_scraping(False)
                    status_pill.configure(text="● Browser: Ready", fg="#4ADE80")

        except queue.Empty:
            pass
        root.after(100, poll_ui_queue)

    def on_close() -> None:
        worker.stop_event.set()
        cmd_q.put(("QUIT", None))
        root.destroy()

    root.protocol("WM_DELETE_WINDOW", on_close)
    append_log(f"Welcome to {APP_TITLE}.")
    append_log(f"Output directory: {OUTPUT_DIR.resolve()} (Only .xlsx files will be saved)")
    root.after(100, poll_ui_queue)
    root.mainloop()
    return True


def prompt_cli_date_range(require_at_least_one: bool = False) -> tuple[date | None, date | None]:
    while True:
        print("\n--- Date Range Filter (Format: DD/MM/YYYY, e.g. 05/10/2026) ---")
        s_raw = input("  From Date (Start Date) [press ENTER for no lower limit]: ").strip()
        e_raw = input("  To Date   (End Date)   [press ENTER for no upper limit]: ").strip()
        if require_at_least_one and not s_raw and not e_raw:
            print("  Please enter at least a From Date or To Date.")
            continue
        s_dt, e_dt, err = normalize_date_range(s_raw, e_raw)
        if err:
            print(f"  [!] {err}")
            continue
        return s_dt, e_dt


def run_cli() -> None:
    """Clean, production-grade Console UI fallback with Date Range support."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    print("=" * 76)
    print(f"  {APP_TITLE}")
    print("=" * 76)
    print(f"  Output Folder : {OUTPUT_DIR.resolve()} (.xlsx only)")
    print("  Row Order     : Lowest Row (Oldest) -> Highest Row (Newest) per page")
    print("  Columns       : DATE | NAME | PRODUCT | COURT | PRICE | TOTAL")
    print("-" * 76)
    print("  [1] Test Lowest Filing on Current Page")
    print("  [2] Scrape Current Page Only (Lowest -> Highest)")
    print("  [3] Reverse Scrape: Current Page -> Page 1 (Lowest -> Highest)")
    print("  [4] Scan Specific Date Range (From Date -> To Date)")
    print("  [5] Exit")
    print("=" * 76)

    choice = input("\nSelect option (1-5): ").strip()
    if choice == "5":
        return
    if choice not in {"1", "2", "3", "4"}:
        print("Invalid selection.")
        return

    start_dt: date | None = None
    end_dt: date | None = None
    if choice == "4":
        start_dt, end_dt = prompt_cli_date_range(require_at_least_one=True)
    elif choice in {"2", "3"}:
        use_filter = input("\nApply a Date Range filter? (y/N): ").strip().lower()
        if use_filter == "y":
            start_dt, end_dt = prompt_cli_date_range(require_at_least_one=False)

    with sync_playwright() as p:
        context, page, temp_prof = launch_edge_resilient(p, log_fn=print)
        try:
            page.goto(URL, wait_until="domcontentloaded", timeout=30000)

            print("\n" + "-" * 76)
            print("  BROWSER READY:")
            print("  1) Log in to EFS manually if prompted.")
            print("  2) Open 'My Filing' and ensure Page Size = 10.")
            print("  3) Navigate to your starting page (or stay on Page 1 for Date Range auto-seek).")
            print("-" * 76)
            input("\nPress ENTER when Microsoft Edge is ready... ")

            try:
                page.wait_for_selector(ROWS_SELECTOR, timeout=20000)
            except PlaywrightTimeoutError:
                print("\nERROR: Could not detect the Filing Grid.")
                return

            page_no = current_page_number(page)
            page_size = current_page_size(page)
            rows = page_rows(page)

            print(f"\nDetected Page: {page_no} | Page Size: {page_size or 'unknown'} | Visible Filings: {len(rows)}")

            if choice in {"3", "4"} and page_size != "10":
                print("\nSTOP: Page Size must be set to 10.")
                return

            if choice == "3" and start_dt is None and end_dt is None and page_no == 1:
                print("\nSTOP: You are on Page 1. Please navigate to an older starting page first, or use Option [4] Date Range.")
                return

            if choice == "1":
                if not rows:
                    print("No filing rows found.")
                else:
                    test_idx = len(rows) - 1
                    filing_id = rows[test_idx]["filing_id"]
                    print(f"\nTesting lowest filing (Row {test_idx + 1}): {filing_id}")
                    result = click_filing_row(page, test_idx)
                    records = apply_date_grouping(result["records"])
                    for r in records:
                        print(
                            f"  {r['DATE']:10} | {r['NAME'][:28]:28} | "
                            f"{r['PRODUCT'][:24]:24} | {r['PRICE']:>8} | {r['TOTAL']:>8}"
                        )
                    test_xlsx = save_test(records)
                    print("\nSAVED EXCEL:", test_xlsx.resolve())
            else:
                # Auto-seek forward to start_dt if in multi-page mode
                if choice in {"3", "4"} and start_dt is not None:
                    while True:
                        p_no = current_page_number(page)
                        cur_rows = page_rows(page)
                        if not cur_rows:
                            break
                        r_dates = [parse_user_date(r["filing_date"]) for r in cur_rows]
                        v_dates = [d for d in r_dates if d is not None]
                        if not v_dates:
                            break
                        if min(v_dates) >= start_dt:
                            print(
                                f"  Auto-seeking to {start_dt.strftime('%d/%m/%Y')}: "
                                f"Page {p_no} oldest is {min(v_dates).strftime('%d/%m/%Y')} -> advancing to Page {p_no + 1}..."
                            )
                            if not click_next_page(page):
                                break
                        else:
                            break

                all_records: list[dict[str, str]] = []
                processed_ids: set[str] = set()
                reached_past_end_date = False

                while not reached_past_end_date:
                    page_no = current_page_number(page)
                    rows = page_rows(page)
                    if not rows:
                        break

                    row_dates = [parse_user_date(r["filing_date"]) for r in rows]
                    valid_dates = [d for d in row_dates if d is not None]

                    if choice in {"3", "4"} and start_dt is not None and valid_dates and max(valid_dates) < start_dt:
                        print(f"  Page {page_no}: All filings older than {start_dt.strftime('%d/%m/%Y')} — skipping page...")
                        if page_no <= 1 or not click_previous_page(page):
                            break
                        continue

                    if choice in {"3", "4"} and end_dt is not None and valid_dates and min(valid_dates) > end_dt:
                        print(f"  Page {page_no}: Filings are newer than {end_dt.strftime('%d/%m/%Y')} — stopping scan.")
                        break

                    print(f"\n--- Page {page_no} ({len(rows)} filings, lowest -> highest) ---")

                    row_order = list(range(len(rows) - 1, -1, -1))
                    for step, i in enumerate(row_order, start=1):
                        filing_id = rows[i]["filing_id"]
                        row_date_str = rows[i].get("filing_date", "")
                        row_dt = parse_user_date(row_date_str)

                        if filing_id in processed_ids:
                            continue

                        if start_dt is not None and row_dt is not None and row_dt < start_dt:
                            print(f"  [Skip] Row {i + 1} ({filing_id}, {row_date_str}) < {start_dt.strftime('%d/%m/%Y')}")
                            continue

                        if end_dt is not None and row_dt is not None and row_dt > end_dt:
                            print(f"  [Stop] Row {i + 1} ({filing_id}, {row_date_str}) > {end_dt.strftime('%d/%m/%Y')}")
                            if choice in {"3", "4"}:
                                reached_past_end_date = True
                                break
                            continue

                        try:
                            started = time.monotonic()
                            result = click_filing_row(page, i)
                            elapsed = time.monotonic() - started
                            processed_ids.add(filing_id)
                            all_records.extend(result["records"])

                            print(
                                f"  [{step}/{len(rows)}] Row {i + 1} ({filing_id}, {row_date_str}) -> "
                                f"{len(result['records'])} row(s) ({elapsed:.2f}s)"
                            )
                            for r in result["records"]:
                                print(
                                    f"       {r['DATE']:10} | {r['NAME'][:25]:25} | "
                                    f"{r['PRODUCT'][:24]:24} | {r['PRICE']:>8} | {r['TOTAL']:>8}"
                                )
                        except Exception as exc:
                            print(f"  [{step}/{len(rows)}] Row {i + 1} ({filing_id}) FAILED: {exc}")

                    if choice == "2" or reached_past_end_date:
                        break

                    if all_records:
                        save_checkpoint(all_records, page_no)

                    if page_no <= 1:
                        break
                    if not click_previous_page(page):
                        print("Could not move to previous page. Stopping safely.")
                        break

                if all_records:
                    suffix = ""
                    if start_dt or end_dt:
                        s_tag = start_dt.strftime("%Y%m%d") if start_dt else "start"
                        e_tag = end_dt.strftime("%Y%m%d") if end_dt else "end"
                        suffix = f"{s_tag}_to_{e_tag}"
                    final_xlsx = save_final(all_records, suffix_tag=suffix)
                    print("\n" + "=" * 76)
                    print(f"  FINAL EXCEL SAVED : {final_xlsx.resolve()}")
                    print(f"  TOTAL ROWS WRITTEN: {len(all_records)}")
                    print("=" * 76)
                else:
                    cleanup_non_xlsx_files()
                    print("\nNo matching document rows were captured in the selected date range.")

            input("\nPress ENTER to close the browser and exit...")
        finally:
            try:
                context.close()
            except Exception:
                pass
            if temp_prof and temp_prof.exists():
                shutil.rmtree(temp_prof, ignore_errors=True)


def run() -> None:
    if "--cli" not in sys.argv:
        if launch_gui():
            return
    run_cli()


if __name__ == "__main__":
    try:
        run()
    except KeyboardInterrupt:
        print("\nStopped by user.")
    except Exception as exc:
        import traceback
        print("\n" + "=" * 76)
        print("  FATAL ERROR:", repr(exc))
        print("=" * 76)
        traceback.print_exc()
        try:
            input("\nPress ENTER to close this window...")
        except Exception:
            time.sleep(15)
        sys.exit(1)